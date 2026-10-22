"""Build the Russian teaching guide from its repository Markdown source.

This formats explanatory text; it never changes measured data, electrical models,
firmware or the authoritative netlists. Run with the bundled document runtime.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "input-explained-ru.md"
OUTPUT = ROOT / "motortester-input-guide-ru.docx"


def link(paragraph, label, target):
    rel = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    el = OxmlElement("w:hyperlink")
    el.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    prop = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "006A76")
    prop.append(color)
    run.append(prop)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    el.append(run)
    paragraph._p.append(el)


def inline(paragraph, value):
    # Preserve emphasis and real external links, without raw Markdown in Word.
    pattern = r"(\*\*.*?\*\*|\[[^\]]+\]\([^)]+\)|https://\S+)"
    for token in re.split(pattern, value):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("["):
            match = re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", token)
            if match:
                link(paragraph, match[1], match[2])
            else:
                paragraph.add_run(token)
        elif token.startswith("https://"):
            # Long datasheet URLs remain clickable without overrunning the page.
            link(paragraph, "Открыть источник", token.rstrip("."))
        else:
            paragraph.add_run(token)


def table(doc, rows):
    cells = [[v.strip() for v in row.strip().strip("|").split("|")] for row in rows]
    cells = [row for row in cells if not all(re.fullmatch(r"[: -]+", v) for v in row)]
    count = len(cells[0])
    # Short numeric/designator columns leave room for electrical destinations.
    if count == 4 and cells[0][0] == "Обозначение":
        widths = [2.8, 1.3, 4.0, 8.4]
    elif count == 3 and cells[0][0] == "Материал":
        widths = [4.5, 2.0, 10.0]
    elif count == 4 and cells[0][0] == "Секция":
        widths = [3.0, 4.2, 4.8, 4.5]
    elif count == 4:
        widths = [2.5, 3.0, 4.6, 6.4]
    elif count == 3:
        widths = [4.1, 5.7, 6.7]
    else:
        widths = [5.7, 10.8]
    t = doc.add_table(rows=0, cols=count)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{side}")
        for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            node.set(qn(f"w:{attr}"), value)
        borders.append(node)
    t._tbl.tblPr.append(borders)
    for index, row in enumerate(cells):
        tcells = t.add_row().cells
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for col, (cell, value) in enumerate(zip(tcells, row)):
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), "85")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            pr.append(margins)
            fill = OxmlElement("w:shd")
            fill.set(qn("w:fill"), "DFEAF0" if index == 0 else ("F5F7F8" if index % 2 == 0 else "FFFFFF"))
            pr.append(fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9.5)
                if index == 0:
                    run.bold = True
            if (count == 4 and col < 2) or (cells[0][0] == "Материал" and col == 1):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.25)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    sec.header_distance = sec.footer_distance = Cm(0.8)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    doc.styles["Title"].font.size = Pt(25)
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    for name, size in (("Heading 1", 17), ("Heading 2", 13)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.italic = False
    doc.styles["Caption"].paragraph_format.space_after = Pt(8)
    hp = sec.header.paragraphs[0]
    hp.text = "Мотортестер V2   •   Учебное руководство по входным схемам"
    for run in hp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.add_run("Редакция 3   •   ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    fp._p.append(field)
    for run in fp.runs:
        run.font.size = Pt(9)
    doc.core_properties.title = "Как работает вход Мотортестера V2"
    doc.core_properties.subject = "Объяснение схем P1-LV, P0, АЦП и будущих работ"
    doc.core_properties.author = "Проект Мотортестер V2"
    doc.core_properties.keywords = "P1-LV, P0, CH1, учебное руководство"
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i, figures = 0, 0
    next_page = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "<!-- pagebreak -->":
            next_page = True
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            p = doc.add_paragraph(line[3:], style="Heading 1")
            if next_page:
                p.paragraph_format.page_break_before = True
                next_page = False
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
            assert match, line
            image = (SOURCE.parent / match[2]).resolve()
            figures += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            width_cm = 13.7 if image.stem == "sallen-key" else 16
            shape = p.add_run().add_picture(str(image), width=Cm(width_cm))
            shape._inline.docPr.set("descr", match[1])
            doc.add_paragraph(f"Рисунок {figures}. {match[1]}.", style="Caption")
        elif line.startswith("|"):
            rows = [line]
            while i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
                i += 1
                rows.append(lines[i].strip())
            table(doc, rows)
        else:
            para = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not lines[i + 1].startswith(("#", "|", "![", "<!--")):
                i += 1
                para.append(lines[i].strip())
            p = doc.add_paragraph()
            inline(p, " ".join(para))
        i += 1
    doc.save(OUTPUT)
    print(f"Created {OUTPUT.name}; {figures} embedded figures; {len(doc.tables)} tables")


if __name__ == "__main__":
    build()
